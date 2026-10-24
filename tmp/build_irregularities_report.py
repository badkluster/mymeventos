from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path('output/Informe_irregularidades_eventos_contratos_2026-10-23.docx')

BLUE = '17365D'
LIGHT_BLUE = 'EAF2F8'
LIGHT_GRAY = 'F5F6F7'
BORDER = 'D9D9D9'
RED = 'C00000'
AMBER = '9C6500'
GREEN = '006100'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=55, start=100, bottom=55, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    tr_pr.append(cant_split)


def set_font(run, size=9, bold=False, color='000000'):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text, *, bold=False, color='000000', size=8.3, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def priority_text(priority):
    return {'Alta': RED, 'Media': AMBER, 'Revisión': GREEN}[priority]


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        set_cell_shading(cell, BLUE)
        write_cell(cell, h, bold=True, color='FFFFFF', size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_index, row in enumerate(rows):
        word_row = table.add_row()
        prevent_row_split(word_row)
        cells = word_row.cells
        for i, value in enumerate(row):
            cells[i].width = Cm(widths[i])
            if r_index % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
            is_priority = headers[i] == 'Prioridad'
            if is_priority:
                write_cell(cells[i], value, bold=True, color=priority_text(value), size=8.15, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                alignment = WD_ALIGN_PARAGRAPH.CENTER if headers[i] in ('Contrato', 'Estado', 'Fecha') else WD_ALIGN_PARAGRAPH.LEFT
                write_cell(cells[i], value, size=8.15, align=alignment)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, size=14 if level == 1 else 11, bold=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, size=9.4, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, size=9.4)
    else:
        r = p.add_run(text)
        set_font(r, size=9.4)
    return p


def money(value):
    return '$ ' + f'{value:,.0f}'.replace(',', '.')


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.35)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.25)
section.right_margin = Cm(1.25)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(9.4)
for style_name in ['Title', 'Heading 1', 'Heading 2']:
    style = styles[style_name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    style.font.color.rgb = RGBColor(0, 0, 0)
# Remove any theme-supplied border from the built-in Title style as well as
# from the concrete title paragraph below.
title_style_ppr = styles['Title']._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn('w:pBdr'))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('M M Eventos  |  Informe interno de regularización  |  Corte 23 de octubre de 2026')
set_font(fr, size=7.5, color='666666')

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(5)
# Word's built-in Title style may carry a colored bottom border in some Office
# themes. The report uses spacing, not a rule, to separate the title block.
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)
tr = title.add_run('Informe de regularizacion de eventos y contratos')
set_font(tr, size=21, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(9)
sr = subtitle.add_run('Acciones pendientes por salón  |  Corte de datos 23 de octubre de 2026')
set_font(sr, size=10.5, color='404040')

add_body(doc, 'El relevamiento identifica 18 contratos que requieren revisión y 5 eventos activos sin un contrato aprobado. La prioridad inmediata es validar los contratos pendientes de aprobación y corregir las discrepancias entre saldo contractual y cuotas abiertas antes de gestionar nuevos cobros.', 'El relevamiento identifica ')
add_body(doc, 'Alcance: 48 eventos y 47 contratos activos o vigentes revisados en la base configurada del sistema. Los registros de catering externo marcados como QA local se presentan en una sección separada; no deben asignarse a las encargadas de salón sin antes confirmar su destino operativo.', 'Alcance: ')

add_heading(doc, 'Resumen de prioridad', 1)
add_table(doc,
          ['Grupo', 'Cantidad', 'Qué requiere', 'Acción inicial'],
          [
              ['Eventos activos sin contrato aprobado', '5', 'El evento no tiene una versión contractual aprobada. Tres casos pertenecen a QA local de catering externo.', 'Completar checklist contractual y aprobar sólo si la documentación está validada.'],
              ['Contratos en borrador o pendientes de aprobación', '7', 'Cinco bloquean la gestión contractual del evento; dos son revisiones V2 que conviven con una versión ya aprobada.', 'Definir si se aprueba, se corrige o se descarta la revisión.'],
              ['Desajustes entre saldo y cuotas abiertas', '10', 'El saldo del contrato no coincide con el saldo remanente del plan, o existe una cuota abierta con saldo contractual cero.', 'Conciliar comprobantes, historial de pagos y plan de cuotas antes de reclamar.'],
              ['Saldo pendiente sin plan de cuotas', '1', 'Hay deuda contractual pero no existe un cronograma para seguimiento o recordatorios.', 'Definir y cargar el plan de pagos acordado.'],
          ],
          [5.0, 1.7, 9.6, 10.3])

add_heading(doc, 'Criterio de clasificación', 1)
add_body(doc, 'Alta: evento activo sin contrato aprobado o saldo pendiente sin plan de cuotas. Media: diferencia entre el saldo contractual y las cuotas abiertas, o cuota abierta con saldo contractual cero. Revisión: versión contractual en borrador cuando el evento ya conserva otra versión aprobada.')
add_body(doc, 'Las encargadas deben verificar el respaldo comercial y los comprobantes antes de modificar cualquier importe, estado o cuota. El objetivo es corregir el registro que no refleje el acuerdo real, no completar datos por estimación.')

salon_sections = {
    'San Carlos': [
        ['Media', 'C-2026-00002', 'Fiesta de 15 Valentina Ríos', 'Valentina Ortiz', 'Confirmado  01/08/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 550.000.', 'Conciliar cobro final y regularizar la cuota residual si el evento está saldado.'],
        ['Media', 'C-2026-00003', 'Fiesta de Egresados Colegio San Martín', 'Valentina Ortiz', 'Confirmado  08/08/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 400.000.', 'Conciliar cobro final y regularizar la cuota residual si el evento está saldado.'],
        ['Media', 'C-2026-00006', 'Cumpleaños infantil Benjamín Torres', 'Victoria Ortiz', 'Reservado  15/08/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 650.000.', 'Revisar comprobante y plan; definir si corresponde cobro o cierre de cuota.'],
        ['Media', 'C-2026-00007', 'Casamiento Florencia Ojeda y Nicolás Herrera', 'Karen Herrera', 'Reservado  15/08/2026', 'Saldo contractual $ 3.500.000; cuotas abiertas $ 4.100.000.', 'Comparar plan con contrato y cobros aplicados. Diferencia a resolver: $ 600.000.'],
        ['Alta', 'C-2026-00011', 'Fiesta de 15 Abril Vargas', 'Martina Sánchez', 'Contrato borrador  05/09/2026', 'Contrato pendiente de aprobación; saldo $ 2.500.000 y 3 cuotas abiertas.', 'Completar revisión contractual y aprobar sólo con documentación validada.'],
        ['Revisión', 'C-2026-00001-V2', 'Cumple de 40 años Manuela', 'Manuela Albarracin', 'Confirmado  16/09/2026', 'Versión V2 en borrador. Existe C-2026-00001 aprobado para el mismo evento.', 'Definir si la revisión V2 debe formalizarse, corregirse o descartarse.'],
        ['Alta', 'C-2026-00042', 'Cumpleaños Manual QA Operador', 'Manual QA Operador', 'Seña pendiente  15/10/2026', 'Saldo $ 4.949.998 sin plan de cuotas.', 'Confirmar acuerdo y cargar cronograma antes de iniciar cobranza o recordatorios.'],
        ['Revisión', 'C-2026-00042-V2', 'Cumpleaños Manual QA Operador', 'Manual QA Operador', 'Seña pendiente  15/10/2026', 'Versión V2 en borrador; cambia el total a $ 6.400.000.', 'Definir si sustituye a la versión aprobada C-2026-00042.'],
    ],
    'Villa Elisa': [
        ['Media', 'C-2026-00004', 'Casamiento Ana Bravo y Martín Coronel', 'Melina Ortiz', 'Reservado  08/08/2026', 'Saldo contractual $ 3.575.000; cuotas abiertas $ 4.075.000.', 'Conciliar plan y cobros aplicados. Diferencia a resolver: $ 500.000.'],
        ['Media', 'C-2026-00010', 'Comunión de Mía Domínguez', 'Fiorella Campos', 'Reservado  29/08/2026', 'Saldo contractual $ 1.910.000; cuotas abiertas $ 2.310.000.', 'Conciliar plan y cobros aplicados. Diferencia a resolver: $ 400.000.'],
        ['Media', 'C-2026-00013', 'Casamiento Paula Reyes y Federico Silva', 'Karen Fernández', 'Confirmado  19/09/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 500.000.', 'Confirmar cobro total y regularizar la cuota residual del plan si corresponde.'],
    ],
    'La Plata': [
        ['Media', 'C-2026-00005', 'Cena anual Constructora del Plata', 'Oscar García', 'Confirmado  08/08/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 1.000.000.', 'Conciliar cobro final y regularizar la cuota residual si el evento está saldado.'],
        ['Media', 'C-2026-00009', 'Casamiento Carla Maldonado y Diego Ferreyra', 'Emma Vargas', 'Confirmado  22/08/2026', 'Saldo $ 0; queda 1 cuota abierta por $ 1.000.000.', 'Conciliar cobro final y regularizar la cuota residual si el evento está saldado.'],
        ['Media', 'C-2026-00012', 'Egresados Instituto La Plata', 'Franco Maldonado', 'Reservado  12/09/2026', 'Saldo contractual $ 8.500.000; cuotas abiertas $ 9.500.000.', 'Conciliar plan y cobros aplicados. Diferencia a resolver: $ 1.000.000.'],
        ['Alta', 'C-2026-00043', 'Cumple de XV de Cami', 'Anabella Stremi', 'Contrato borrador  09/10/2026', 'Contrato pendiente de aprobación; saldo $ 3.600.000 y 9 cuotas abiertas.', 'Completar revisión contractual y aprobar sólo con documentación validada.'],
    ],
}

for salon, rows in salon_sections.items():
    # Keep La Plata together on a fresh page. Without this break, the last two
    # rows split from the table and Word can place their left edge outside the
    # printable area in the following page.
    if salon == 'La Plata':
        doc.add_page_break()
    add_heading(doc, salon, 1)
    add_table(doc,
              ['Prioridad', 'Contrato', 'Evento', 'Cliente', 'Estado y fecha', 'Situación detectada', 'Acción para la encargada'],
              rows,
              [1.5, 2.5, 4.5, 2.8, 2.6, 5.7, 6.8])

add_heading(doc, 'Catering externo QA local', 1)
add_body(doc, 'Estos tres casos no poseen salón asociado y están identificados como QA local. Se incluyen para trazabilidad, pero no deben asignarse a las encargadas de San Carlos, Villa Elisa o La Plata hasta confirmar si son datos de prueba o servicios reales de catering externo.')
add_table(doc,
          ['Prioridad', 'Contrato', 'Evento', 'Cliente', 'Fecha', 'Situación detectada', 'Acción responsable'],
          [
              ['Alta', 'C-2026-00044', 'Catering Corporativo Parque Norte', 'Cliente Catering 1', '12/11/2026', 'Contrato pendiente de aprobación; saldo $ 1.200.000.', 'Administración: validar si es QA o real; luego aprobar, corregir o desestimar.'],
              ['Alta', 'C-2026-00045', 'Catering Aniversario Terraza Centro', 'Cliente Catering 2', '19/11/2026', 'Contrato pendiente de aprobación; saldo $ 880.000.', 'Administración: validar si es QA o real; luego aprobar, corregir o desestimar.'],
              ['Alta', 'C-2026-00046', 'Catering Graduación Club del Lago', 'Cliente Catering 3', '03/12/2026', 'Contrato pendiente de aprobación; saldo $ 1.850.000.', 'Administración: validar si es QA o real; luego aprobar, corregir o desestimar.'],
          ],
          [1.5, 2.5, 4.8, 3.0, 2.1, 5.4, 6.6])

add_heading(doc, 'Orden sugerido de resolución', 1)
add_table(doc,
          ['Orden', 'Responsable', 'Acción', 'Resultado esperado'],
          [
              ['1', 'Encargada del salón', 'Revisar los eventos sin contrato aprobado y completar el checklist contractual.', 'Cada evento operativo queda respaldado por una versión contractual válida.'],
              ['2', 'Encargada y administración', 'Conciliar contrato, comprobantes y plan de cuotas en los 10 desajustes financieros.', 'Saldo contractual y cuotas abiertas muestran el mismo pendiente real.'],
              ['3', 'Administración', 'Definir el cronograma de C-2026-00042 antes de iniciar la cobranza.', 'El saldo pendiente puede seguirse por vencimientos y recordatorios.'],
              ['4', 'Administración', 'Clasificar los tres registros QA local de catering externo.', 'No quedan tareas operativas asignadas a un salón por error.'],
              ['5', 'Administración', 'Decidir el destino de las revisiones V2 de C-2026-00001 y C-2026-00042.', 'Sólo una versión contractual refleja el acuerdo vigente por evento.'],
          ],
          [1.5, 4.0, 11.0, 12.0])

doc.core_properties.title = 'Informe de regularizacion de eventos y contratos'
doc.core_properties.subject = 'Acciones pendientes por salón'
doc.core_properties.author = 'M M Eventos'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
